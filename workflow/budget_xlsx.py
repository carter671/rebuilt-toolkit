"""takeoff.json (from compute.py) -> rehab budget workbook with live formulas.

    python3 build_xlsx.py takeoff.json OUT.xlsx [--prices PRICEBOOK.json]
then:  python3 <xlsx skill>/scripts/recalc.py OUT.xlsx 60     (must report 0 errors)

Unit costs: the Rebuilt pricebook (rebuilt/takeoff/pricebook.py, or --prices), plus ADDED rates below that the
pricebook does not have yet (shown in red, "added — verify"). Quantities are written as formulas that point at
the takeoff tabs, so changing a room / opening row updates the budget.
"""
import argparse, json, math, sys
import os, pathlib; sys.path.insert(0, os.environ.get('REBUILT_SRC', str(pathlib.Path(__file__).resolve().parents[1] / 'src')))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from rebuilt.takeoff.pricebook import load_pricebook

ap = argparse.ArgumentParser()
ap.add_argument('takeoff'); ap.add_argument('out'); ap.add_argument('--prices', default=None)
A = ap.parse_args()
t = json.load(open(A.takeoff))
PB = load_pricebook(A.prices)

FONT = 'Arial'
f = lambda **k: Font(name=FONT, size=k.pop('size', 10), **k)
BLUE, GREEN, RED = '0000FF', '008000', 'C0392B'
HEAD = PatternFill('solid', fgColor='1F2A36'); SUB = PatternFill('solid', fgColor='E8ECF0'); YEL = PatternFill('solid', fgColor='FFFF00')
MONEY, MONEY2 = '$#,##0;($#,##0);-', '$#,##0.00;($#,##0.00);-'
NUM, NUM1 = '#,##0;(#,##0);-', '#,##0.0;(#,##0.0);-'
wb = Workbook()


def sheet(name, title, sub, widths):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = f(bold=True, size=14)
    ws['A2'] = sub; ws['A2'].font = f(italic=True, color='555555')
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def header(ws, row, cols, freeze=True):
    for i, c in enumerate(cols):
        if c == '':
            continue
        x = ws.cell(row, i + 1, c); x.font = f(bold=True, color='FFFFFF'); x.fill = HEAD
        x.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 30
    if freeze:
        ws.freeze_panes = ws.cell(row + 1, 1)


def subhead(ws, row, cols):
    for i, c in enumerate(cols):
        if c:
            x = ws.cell(row, i + 1, c); x.font = f(bold=True); x.fill = SUB


def put(ws, r, c, v, fmt=None, color=None, bold=False, fill=None):
    x = ws.cell(r, c, v)
    x.font = f(color=color, bold=bold) if color else f(bold=bold)
    if fmt: x.number_format = fmt
    if fill: x.fill = fill
    return x


GAR = lambda n: 'garage' in n.lower()
rooms = sorted(t['rooms'], key=lambda r: (GAR(r['name']), r['storey'] != '1F', r['name']))
room_names = [r['name'] for r in rooms]
baths = [n for n in room_names if 'bath' in n.lower()]
kitchens = [n for n in room_names if 'kitchen' in n.lower()]
laundry = [n for n in room_names if 'laundry' in n.lower() or 'utility' in n.lower()]
bedrooms = [n for n in room_names if 'bed' in n.lower() and 'closet' not in n.lower()]

# ======================================================================= Rooms
R = sheet('Rooms', 'Rooms — areas from the model', 'Floor area to finished faces; wall area net of doors/windows; vaulted ceilings along the slope. Garages are left unfinished (Finished = N).',
          [30, 8, 14, 11, 12, 14, 12, 9, 11, 10, 10])
header(R, 4, ['Room', 'Level', 'Floor finish', 'Floor sf', 'Wall sf (net)', 'Ceiling', 'Ceiling sf', 'Wall ht (ft)', 'Perimeter lf', 'Base lf', 'Finished? (Y/N)'])
r0 = 5
for i, rm in enumerate(rooms):
    r = r0 + i
    put(R, r, 1, rm['name']); put(R, r, 2, rm['storey']); put(R, r, 3, rm['floor'] or 'none', color=BLUE)
    for c, k in ((4, 'floor_sf'), (5, 'wall_sf'), (7, 'ceiling_sf'), (9, 'perim_lf'), (10, 'base_lf')):
        put(R, r, c, round(rm[k], 1), NUM1, BLUE)
    put(R, r, 6, rm['ceiling']); put(R, r, 8, round(rm['height_ft'], 2), '0.00', BLUE); put(R, r, 11, 'Y' if rm['finished'] else 'N', color=BLUE)
rN = r0 + len(rooms) - 1
RT = rN + 2
put(R, RT, 1, 'Totals — finished rooms', bold=True)
for c in (4, 5, 7, 9, 10):
    L = get_column_letter(c); put(R, RT, c, f'=SUMIFS({L}{r0}:{L}{rN},$K${r0}:$K${rN},"Y")', NUM, bold=True)
RNG = lambda col: f"Rooms!${col}${r0}:${col}${rN}"
FIN = lambda col: f'SUMIFS({RNG(col)},{RNG("K")},"Y")'

# ======================================================================= Openings
O = sheet('Openings', 'Window + door schedule', 'Every opening in the model. Action drives the budget — change it here: New / Replace / Reinstall / Keep / Add door / Cased / Bifold / Leave open.',
          [12, 15, 44, 34, 7, 7, 7, 10, 13, 9, 10])
header(O, 4, ['ID', 'Type', 'Name', 'Room(s)', 'W (in)', 'H (in)', 'Sill (in)', 'New / existing', 'Action', 'Area sf', 'Cut into existing wall'])
ops = sorted(t['openings'], key=lambda o: (o['type'], o['name']))
o0 = 5
for i, o in enumerate(ops):
    r = o0 + i
    put(O, r, 1, o['id']); put(O, r, 2, o['type']); put(O, r, 3, o['name']); put(O, r, 4, o['rooms'])
    put(O, r, 5, o['w_in'], NUM, BLUE); put(O, r, 6, o['h_in'], NUM, BLUE); put(O, r, 7, o['sill_in'], NUM, BLUE)
    put(O, r, 8, 'New' if o['new'] else 'Existing'); put(O, r, 9, o['action'], color=BLUE)
    put(O, r, 10, f'=E{r}*F{r}/144', NUM1); put(O, r, 11, 'Y' if o['cut_in_existing_wall'] else '', color=BLUE)
oN = o0 + len(ops) - 1
OC = lambda ty, act: f'COUNTIFS(Openings!$B${o0}:$B${oN},"{ty}",Openings!$I${o0}:$I${oN},"{act}")'

# ======================================================================= Cabinets & fixtures
K = sheet('Cabinets & Fixtures', 'Cabinets, vanities, plumbing fixtures + appliances', 'Placed = drawn in the design. Allowance = needed but not drawn yet (edit to suit).',
          [26, 28, 7, 7, 7, 6, 10, 16, 30])
header(K, 4, ['Room', 'Item', 'W (in)', 'D (in)', 'H (in)', 'Qty', 'Linear ft', 'Category', 'Source'])
CAT = {'base cabinet': 'Base cabinet', 'sink base cabinet': 'Base cabinet', 'kitchen island': 'Island', 'wall cabinet': 'Upper cabinet',
       'tall cabinet': 'Tall cabinet', 'vanity': 'Vanity', 'bathtub': 'Tub', 'shower': 'Shower', 'toilet': 'Toilet',
       'washer': 'Laundry', 'dryer': 'Laundry', 'range': 'Appliance', 'refrigerator': 'Appliance', 'dishwasher': 'Appliance'}
rows = []
for it in sorted(t['items'], key=lambda x: (x['room'], x['cat'])):
    rows.append((it['room'], it['name'] + (' (sink base)' if it['cat'] == 'sink base cabinet' else ''), it['w_in'], it['d_in'], it['h_in'], 1, CAT.get(it['cat'], it['cat']), 'Placed in design'))
for rm in t['existing_toilets']:
    rows.append((rm, 'Toilet (in the model)', 20, 28, 30, 1, 'Toilet', 'In the model'))
have = lambda room, cat: any(r[0] == room and r[6] == cat for r in rows)
for b in baths:                      # every bath gets a vanity + toilet
    half = 'half' in b.lower() or 'powder' in b.lower()
    if not have(b, 'Vanity'):
        w = 24 if half else 36
        rows.append((b, f'Vanity {w}" with top', w, 21, 34, 1, 'Vanity', 'Allowance — not drawn yet'))
    if not have(b, 'Toilet'):
        rows.append((b, 'Toilet', 20, 28, 30, 1, 'Toilet', 'Allowance — not drawn yet'))
for k in kitchens:
    for nm, w, cat in (('Range 30"', 30, 'Appliance'), ('Refrigerator 36"', 36, 'Appliance'), ('Dishwasher 24"', 24, 'Appliance'), ('Range hood / micro', 30, 'Appliance')):
        if not any(r[0] == k and nm.split()[0].lower() in r[1].lower() for r in rows):
            rows.append((k, nm, w, 26, 34, 1, cat, 'Allowance — not drawn yet'))
k0 = 5
for i, (rm, nm, w, d, h, q, cat, src) in enumerate(rows):
    r = k0 + i
    put(K, r, 1, rm); put(K, r, 2, nm); put(K, r, 3, w, NUM, BLUE); put(K, r, 4, d, NUM, BLUE); put(K, r, 5, h, NUM, BLUE)
    put(K, r, 6, q, NUM, BLUE); put(K, r, 7, f'=C{r}/12*F{r}', NUM1); put(K, r, 8, cat); put(K, r, 9, src)
    if src.startswith('Allowance'):
        K.cell(r, 9).font = f(color=RED)
kN = k0 + len(rows) - 1
KS = kN + 2
KR = lambda col: f"'Cabinets & Fixtures'!${col}${k0}:${col}${kN}"
put(K, KS, 1, 'Countertops', bold=True)
put(K, KS + 1, 1, 'Base runs (25.5" deep)'); put(K, KS + 1, 7, f'=SUMIFS(G{k0}:G{kN},H{k0}:H{kN},"Base cabinet")*25.5/12', NUM1); put(K, KS + 1, 8, 'sf')
put(K, KS + 2, 1, 'Island (as drawn)'); put(K, KS + 2, 7, f'=SUMPRODUCT((H{k0}:H{kN}="Island")*C{k0}:C{kN}*D{k0}:D{kN})/144', NUM1); put(K, KS + 2, 8, 'sf')
put(K, KS + 3, 1, 'Total quartz', bold=True); put(K, KS + 3, 7, f'=G{KS+1}+G{KS+2}', NUM1, bold=True); put(K, KS + 3, 8, 'sf')
QUARTZ = f"'Cabinets & Fixtures'!$G${KS+3}"
KLF = lambda cat: f'SUMIFS({KR("G")},{KR("H")},"{cat}")'
KCNT = lambda cat: f'SUMIFS({KR("F")},{KR("H")},"{cat}")'
if kitchens and not any(r[6] == 'Upper cabinet' and r[0] in kitchens for r in rows):
    put(K, KS + 5, 1, 'The kitchen has no upper cabinets drawn — add them in the editor if you want them.').font = f(italic=True, color=RED)
# vanities by size, lavs (a 60"+ vanity has two sinks), tile surrounds, showers
van = [r for r in rows if r[6] == 'Vanity']
lavs = sum(2 if r[2] >= 60 else 1 for r in van)
tile_wall = sum((r[2] + 2 * r[3]) * 72 / 144 for r in rows if r[6] == 'Tub') + sum((r[2] + 2 * r[3]) * 96 / 144 for r in rows if r[6] == 'Shower')
n_show = sum(1 for r in rows if r[6] == 'Shower')

# ======================================================================= Framing & structure
S = sheet('Framing & Structure', 'Framing + structure', 'New / rebuilt wall framing from the model (wall sf includes gable triangles, net of openings). Studs @16" o.c. + 4 per opening + 2 per wall; plates = 3x length.',
          [52, 7, 15, 10, 9, 10, 9, 9, 9, 10, 11])
header(S, 4, ['Wall', 'Level', 'Type', 'Status', 'New lf', 'Wall sf', 'Avg ht', 'Openings', 'Studs', 'Plates lf', 'Sheathing 4x8'])
s0 = 5
fr = sorted(t['framing'], key=lambda w: (w['type'], w['storey'], -w['lf']))
for i, w in enumerate(fr):
    r = s0 + i
    nm = w['wall'] if w['wall'] not in ('Addition wall', 'New wall', 'New outside wall') else f"{w['wall']} ({w['id']})"
    put(S, r, 1, nm); put(S, r, 2, w['storey']); put(S, r, 3, w['type']); put(S, r, 4, w['status'])
    put(S, r, 5, round(w['lf'], 1), NUM1, BLUE); put(S, r, 6, round(w['sf'], 1), NUM1, BLUE); put(S, r, 7, f'=IFERROR(F{r}/E{r},0)', '0.0')
    put(S, r, 8, w['openings'], NUM, BLUE); put(S, r, 9, f'=ROUNDUP(E{r}*12/16,0)+1+4*H{r}+2', NUM)
    put(S, r, 10, f'=3*E{r}', NUM1); put(S, r, 11, f'=IF(LEFT(C{r},3)="2x6",ROUNDUP(F{r}/32*1.1,0),0)', NUM)
sN = s0 + max(len(fr), 1) - 1
put(S, sN + 1, 1, 'Totals', bold=True)
for c, fm in ((5, NUM1), (6, NUM), (9, NUM), (10, NUM), (11, NUM)):
    L = get_column_letter(c); put(S, sN + 1, c, f'=SUM({L}{s0}:{L}{sN})', fm, bold=True)
FR = lambda col, ty: f"SUMIFS('Framing & Structure'!${col}${s0}:${col}${sN},'Framing & Structure'!$C${s0}:$C${sN},\"{ty}\")"
SR = lambda col, a, b: f"'Framing & Structure'!${col}${a}:${col}${b}"


def block(title, cols, data, after):
    """a titled table under the last one; returns (first, last, row after)"""
    r = after + 2
    put(S, r, 1, title, bold=True); r += 1
    subhead(S, r, cols); r += 1
    a = r
    for row in data:
        for c, v in row.items():
            put(S, r, c, v[0], v[1] if len(v) > 1 else None, v[2] if len(v) > 2 else None)
        r += 1
    return a, max(a, r - 1), r


ia, ib, nxt = block('Outside walls to insulate (a finished room on one side only)', ['Wall', 'Level', 'Next to', '', '', 'Wall sf'],
                    [{1: (w['wall'],), 2: (w['storey'],), 3: (w['rooms'],), 6: (round(w['sf']), NUM, BLUE)} for w in t['insulate']] or [{1: ('none',)}], sN + 1)
put(S, nxt, 1, 'Total', bold=True); put(S, nxt, 6, f'=SUM(F{ia}:F{ib})', NUM, bold=True); INS_EXT = f"'Framing & Structure'!$F${nxt}"
fa, fb, nxt = block('Fire-rated walls (garage beside a finished room): 5/8" Type X on the garage side', ['Wall', '', '', '', '', 'Wall sf'],
                    [{1: (w['wall'],), 6: (round(w['sf']), NUM, BLUE)} for w in t['firewall']] or [{1: ('none',), 6: (0, NUM, BLUE)}], nxt)
put(S, nxt, 1, 'Total', bold=True); put(S, nxt, 6, f'=SUM(F{fa}:F{fb})', NUM, bold=True); FIRE = f"'Framing & Structure'!$F${nxt}"
demo = [{1: (d['name'],), 3: ('interior',), 5: (round(d['len_ft'], 1), NUM1, BLUE)} for d in t['demo_int']] + \
       [{1: (d['name'],), 3: ('exterior',), 5: (round(d['len_ft'], 1), NUM1, BLUE)} for d in t['demo_ext']]
da, db, nxt = block('Walls taken out', ['Wall', '', 'Side', '', 'Length ft'], demo or [{1: ('none',), 3: ('interior',), 5: (0, NUM1, BLUE)}], nxt)
DEMO = lambda side: f"SUMIFS({SR('E', da, db)},{SR('C', da, db)},\"{side}\")"
beams = t['beams']
STR = t.get('structure') or {}
for b in beams:
    b.setdefault('plies', 3); b.setdefault('size', f"3-ply 1-3/4\" x {b['depth_in']:g}\" LVL")
ba, bb, nxt = block('Beams + posts (first-pass sizes, rebuilt.engineering — engineer to confirm)', ['Beam', '', 'Size', '', 'Length ft', 'Depth in', 'Spans', 'Plies', 'Load each post lb', 'Post + footing'],
                    [{1: (b['name'],), 3: (b['size'],),
                      5: (round(b['len_ft'], 1), NUM1, BLUE), 6: (b['depth_in'], NUM1, BLUE), 7: (b['count'], NUM, BLUE), 8: (b['plies'], NUM, BLUE),
                      9: (b.get('reaction_lb', 0), NUM, BLUE), 10: (f"{b.get('post', '')} on {b.get('pad', '')}" if b.get('post') else 'engineer',)} for b in beams] or [{1: ('none',), 7: (0, NUM, BLUE)}], nxt)
if STR.get('headers'):
    ha, hb, nxt = block('Headers over new openings in outside walls (first pass)', ['Opening', '', 'Header', '', 'Width in', 'Load plf'],
                        [{1: (h['name'],), 3: (h['size'],), 5: (h['w_in'], NUM, BLUE), 6: (h['load_plf'], NUM, BLUE)} for h in STR['headers']], nxt)
# roof framing for every new roof
rfr = []
for rf in t['roofs']:
    if not rf['new']:
        continue
    span = rf['run_ft'] * (2 if rf['form'] == 'gable' else 1)
    sized = next((x for x in STR.get('rafters', []) if x['roof'] == rf['id']), None)
    if (sized and sized['size'].startswith('engineered') and rf['form'] == 'gable') or (not sized and rf['form'] == 'gable' and span > 24):
        rfr.append((f"{rf['name']}", 'Engineered trusses @24" (2 gable-end)', math.ceil(rf['length_ft'] / 2) + 1, round(span, 1)))
    else:
        size = sized['size'].split()[0] if sized and not sized['size'].startswith('engineered') else ('2x12' if sized else ('2x10' if rf['run_ft'] > 12 else '2x8'))
        n = (math.ceil(rf['length_ft'] * 12 / 16) + 1) * (2 if rf['form'] == 'gable' else 1)
        rfr.append((f"{rf['name']}", f'{size} rafters @16"' + (f" ({sized['span_ft']:g} ft clear span)" if sized else '') + (' — engineer: past the 2x12 span' if sized and sized['size'].startswith('engineered') else ''), n,
                    round(rf['run_ft'] * math.sqrt(1 + (rf['pitch_in12'] / 12) ** 2) + 1, 1)))
ra, rb, nxt = block('Roof framing (new roofs)', ['Roof', '', 'Members', '', 'Count', 'Length ft'],
                    [{1: (a,), 3: (b,), 5: (c, NUM, BLUE), 6: (d, NUM1, BLUE)} for a, b, c, d in rfr] or [{1: ('none',), 5: (0, NUM, BLUE), 6: (0, NUM1, BLUE)}], nxt)

# ======================================================================= Roof & exterior
X = sheet('Roof & Exterior', 'Roofing, gutters + siding', 'Roof areas are the finished roof surfaces in the model after tie-ins are cut (1 square = 100 sf). Eaves/rakes count only the stretches still exposed.',
          [60, 9, 11, 10, 10, 10, 10, 26])
header(X, 4, ['Roof', 'Pitch', 'Roof sf', 'Squares', 'Eave lf', 'Rake lf', 'Ridge lf', 'Action'])
x0 = 5
roofs = sorted(t['roofs'], key=lambda r: (r['new'], r['id']))
for i, rf in enumerate(roofs):
    r = x0 + i
    put(X, r, 1, rf['name']); put(X, r, 2, f"{rf['pitch_in12']:.1f}/12"); put(X, r, 3, round(rf['sf']), NUM, BLUE)
    put(X, r, 4, f'=C{r}/100', '0.0'); put(X, r, 5, round(rf['eave_lf'], 1), NUM1, BLUE); put(X, r, 6, round(rf['rake_lf'], 1), NUM1, BLUE)
    put(X, r, 7, round(rf['ridge_lf'], 1), NUM1, BLUE); put(X, r, 8, 'New framing + roof' if rf['new'] else 'Tear off + re-roof to match', color=BLUE)
xN = x0 + len(roofs) - 1
put(X, xN + 1, 1, 'Totals', bold=True)
for c, fm in ((3, NUM), (4, '0.0'), (5, NUM1), (6, NUM1), (7, NUM1)):
    L = get_column_letter(c); put(X, xN + 1, c, f'=SUM({L}{x0}:{L}{xN})', fm, bold=True)
XR = lambda col: f"'Roof & Exterior'!${col}${x0}:${col}${xN}"
ROOF_EXIST = f'SUMIFS({XR("C")},{XR("H")},"Tear off*")'
ROOF_NEW = f'SUMIFS({XR("C")},{XR("H")},"New*")'
low = [x0 + i for i, rf in enumerate(roofs) if rf['pitch_in12'] < 3]
ROOF_LOW = '+'.join(f"'Roof & Exterior'!$C${r}" for r in low) or '0'
E = t['edges']
r = xN + 4
put(X, r, 1, 'Edges, flashing + water', bold=True); r += 1
subhead(X, r, ['Item', '', 'Qty', 'Unit']); r += 1
EDGE = {}
for nm, q, u, key in [('Gutters — every exposed eave', round(E['gutter_lf']), 'lf', 'gut'),
                      ('Downspouts (~1 per 35 ft of gutter, 10 ft each)', E['downspout_lf'], 'lf', 'ds'),
                      ('Drip edge — eaves + rakes', f'=E{xN+1}+F{xN+1}', 'lf', 'drip'),
                      ('Ridge cap', f'=G{xN+1}', 'lf', 'ridge'),
                      ('Valleys (tie-in gables)', round(E['valley_lf']), 'lf', 'valley'),
                      ('Roof-to-wall / transition flashing (new roofs)', round(E['step_lf']), 'lf', 'step'),
                      ('Ice & water shield — whole low-slope roofs (< 3/12)', f'={ROOF_LOW}', 'sf', 'iw1'),
                      ('Ice & water shield — 3 ft at the other eaves + valleys', '', 'sf', 'iw2'),
                      ('Synthetic underlayment — everything else', '', 'sf', 'syn'),
                      ('Siding + trim — new outside walls (net of openings)', '=' + FR('F', '2x6 exterior'), 'sf', 'sid'),
                      ('Soffit — new eaves + rakes', round(E['soffit_lf']), 'lf', 'sof'),
                      ('Fascia — new eaves + rakes', round(E['soffit_lf']), 'lf', 'fas')]:
    put(X, r, 1, nm); put(X, r, 3, q, NUM, None if str(q).startswith('=') else BLUE); put(X, r, 4, u)
    EDGE[key] = f"'Roof & Exterior'!$C${r}"; r += 1
steep_eaves = '+'.join(f'E{x0 + i}' for i, rf in enumerate(roofs) if rf['pitch_in12'] >= 3) or '0'
X[EDGE['iw2'].split('!')[1].replace('$', '')] = f"=({steep_eaves})*3+{EDGE['valley'].split('!')[1]}*3"
X[EDGE['syn'].split('!')[1].replace('$', '')] = f"=MAX(0,C{xN+1}-{EDGE['iw1'].split('!')[1]})"
for k in ('iw2', 'syn'):
    X[EDGE[k].split('!')[1].replace('$', '')].number_format = NUM

# ======================================================================= Budget
B = sheet('Budget', f"Rehab budget — {t['name']}", 'Qty pulls from the takeoff tabs (black = formula, blue = typed). Unit costs are installed (labor + material). Include: 1 = in, 0 = out.',
          [18, 48, 46, 11, 7, 11, 8, 13, 26])
header(B, 4, ['Trade', 'Item', 'Where / notes', 'Qty', 'Unit', 'Unit cost', 'Include', 'Total', 'Rate source'])
ADDED = {
    'cab_base_only': ('Base cabinets (no uppers above)', 'lf', 300.0), 'cab_upper': ('Upper cabinets', 'lf', 180.0),
    'vanity_72': ('Double vanity 60"+ with quartz top + 2 sinks', 'ea', 2400.0), 'vanity_36': ('30-48" vanity + top + sink', 'ea', 1100.0),
    'vanity_24': ('Vanity under 30" + top + sink', 'ea', 750.0), 'tile_wall': ('Wall tile installed (tub / shower surrounds)', 'sf', 16.0),
    'shower_pan': ('Tiled shower pan + curb + waterproofing', 'ea', 1200.0), 'shower_glass': ('Frameless shower door / panel', 'ea', 1400.0),
    'toilet': ('Toilet (fixture)', 'ea', 350.0), 'tub': ('Alcove tub (fixture)', 'ea', 650.0), 'shower_valve': ('Tub / shower valve + trim', 'ea', 450.0),
    'lav_faucet': ('Lav faucet + drain', 'ea', 180.0), 'kitchen_sink': ('Kitchen sink + faucet + disposal', 'ea', 650.0),
    'bath_access': ('Mirror, bars, paper holder, towel ring', 'ea', 350.0), 'door_garage_reset': ('Remove + reinstall an existing overhead door', 'ea', 900.0),
    'ice_water': ('Ice & water shield (extra over underlayment)', 'sf', 1.25), 'flashing': ('Valley / step / transition flashing', 'lf', 12.0),
    'recessed': ('Recessed LED light', 'ea', 165.0), 'firewall': ('5/8" Type X drywall, garage side', 'sf', 3.75),
    'insul_spray': ('Closed-cell spray foam at vaulted ceilings (R-38+)', 'sf', 4.50), 'turf': ('Golf sim turf / carpet', 'sf', 6.00),
    'drip': ('Drip edge', 'lf', 3.50), 'ridge_cap': ('Ridge cap + ridge vent', 'lf', 9.00),
    'cab_bumpout': ('Framed bump-out behind cabinets set flush with the fridge (2x4 + drywall)', 'lf', 45.0),
    'trash_pullout': ('Trash pull-out kit (double bin) in a base cabinet', 'ea', 220.0), 'lazy_susan': ('Lazy susan in a corner base', 'ea', 260.0),
}
for k, v in PB.items():          # a real pricebook entry always wins over an added rate
    ADDED.pop(k, None)


def rate(k):
    if k in ADDED:
        d, u, p = ADDED[k]; return d, u, p, 'added — verify'
    v = PB[k]
    src = v.get('source', '')
    return v['label'], v['unit'], v['cost'], f"pricebook{' (placeholder)' if src == 'placeholder' else ''}: {k}"


lines = []
L = lambda *a: lines.append(a)
fin_sf = FIN('D')
added_sf = round(sum(s['sf'] for s in t['slabs'] if not s.get('stretched')))
L('General', 'Building permits', '', 1, 'permit_addition' if added_sf else 'permit', 1)
L('General', 'Permit drawings / design professional', 'from the model', 1, 'permit_set', 1 if (added_sf or beams) else 0)
L('General', 'Structural engineer', 'beams, footings, roof changes', 1, 'engineer', 1 if (beams or added_sf) else 0)
L('General', 'Dumpsters (20 yd)', '~1 per 500 finished sf on a gut', f'=MAX(2,ROUNDUP({fin_sf}/500,0))', 'dumpster', 1)
L('General', 'Final clean', 'finished floor area', f'={fin_sf}', 'clean', 1)
L('Demo', 'Remove interior walls', 'Framing sheet: walls taken out', f'={DEMO("interior")}', 'wall_remove', 1)
L('Demo', 'Remove outside wall sections', 'headers priced under Structure', f'={DEMO("exterior")}', 'wall_remove', 1)
if t['roof_strip_sf']:
    L('Demo', 'Strip the old roof slope where the new roof ties in', 'shingles, sheathing + rafters', round(t['roof_strip_sf']), 'demo_select', 1)
for s in t['slabs']:
    L('Foundation', s['name'].replace(' floor (addition)', ' — addition foundation'), 'floor level with the house' if not s.get('stretched') else '', round(s['sf']), 'foundation', 1)
L('Framing', 'Exterior wall framing, 2x6, sheathing + wrap', 'Framing sheet', f"={FR('F', '2x6 exterior')}", 'frame_ext', 1)
L('Framing', 'Interior partitions, 2x4', 'Framing sheet', f"={FR('F', '2x4 partition')}", 'frame_int', 1)
L('Framing', 'Roof framing + sheathing — new roofs', 'Roof sheet', f'={ROOF_NEW}', 'frame_roof', 1)
for i, b in enumerate(beams):
    L('Structure', b['name'], f"{round(b['len_ft'], 1)} ft, {b['size']}" + (f", {b['count']} spans" if b['count'] > 1 else ''), f"='Framing & Structure'!G{ba + i}", 'beam', 1)
L('Roofing', 'Tear-off + re-roof existing roof to match', 'Roof sheet: action "Tear off"', f'={ROOF_EXIST}', 'roof_replace', 1)
L('Roofing', 'Shingles — new roofs', 'Roof sheet: action "New"', f'={ROOF_NEW}', 'roof_replace', 1)
L('Roofing', 'Ice & water shield', 'low-slope roofs + eaves + valleys', f"={EDGE['iw1']}+{EDGE['iw2']}", 'ice_water', 1)
L('Roofing', 'Drip edge', '', f"={EDGE['drip']}", 'drip', 1)
L('Roofing', 'Ridge cap + ridge vent', '', f"={EDGE['ridge']}", 'ridge_cap', 1)
L('Roofing', 'Valley + roof-to-wall flashing', '', f"={EDGE['valley']}+{EDGE['step']}", 'flashing', 1)
L('Exterior', 'Siding + trim on new outside walls', 'to match existing', f"={EDGE['sid']}", 'siding_new', 1)
L('Exterior', 'Soffit — new eaves + rakes', '', f"={EDGE['sof']}", 'soffit_vinyl', 1)
L('Exterior', 'Fascia wrap — new eaves + rakes', '', f"={EDGE['fas']}", 'fascia_wrap', 1)
L('Exterior', 'Seamless gutters', '', f"={EDGE['gut']}", 'gutter_new', 1)
L('Exterior', 'Downspouts', '', f"={EDGE['ds']}", 'downspout_replace', 1)
egress = f'COUNTIFS(Openings!$B${o0}:$B${oN},"Window",Openings!$I${o0}:$I${oN},"New",Openings!$C${o0}:$C${oN},"*egress*")'
L('Windows & Doors', 'New windows', 'Openings: Window / New', f'={OC("Window", "New")}-{egress}', 'window_std', 1)
L('Windows & Doors', 'Egress windows', 'Openings: name contains "egress"', f'={egress}', 'win_egress', 1)
L('Windows & Doors', 'Replace existing windows', 'Openings: Replace — change to Keep to save them', f'={OC("Window", "Replace")}', 'window_std', 1)
L('Windows & Doors', 'New openings cut into existing walls', 'Openings: Cut into existing wall = Y', f'=COUNTIFS(Openings!$K${o0}:$K${oN},"Y")', 'opening_new', 1)
L('Windows & Doors', 'Exterior doors', 'Openings: Exterior door / New', f'={OC("Exterior door", "New")}', 'door_ext', 1)
L('Windows & Doors', 'New garage overhead doors', 'Openings: Garage door / New', f'={OC("Garage door", "New")}', 'door_garage', 1)
L('Windows & Doors', 'Reinstall existing garage doors', 'Openings: Garage door / Reinstall', f'={OC("Garage door", "Reinstall")}', 'door_garage_reset', 1)
L('Windows & Doors', 'Interior prehung doors', 'new doors + doors added to open doorways', f'={OC("Interior door", "New")}+{OC("Doorway", "Add door")}', 'door_int', 1)
L('Windows & Doors', 'Cased openings', 'Openings: Cased', f'={OC("Doorway", "Cased")}', 'cased_opening', 1)
L('Windows & Doors', 'Closet bifolds', 'Openings: Bifold', f'={OC("Closet opening", "Bifold")}', 'door_bifold', 1)
L('Insulation', 'Wall insulation R-20 — outside walls', 'Framing sheet: walls to insulate', f'={INS_EXT}', 'insul_wall', 1)
L('Insulation', 'Blown insulation R-49 — flat ceilings', 'finished rooms', f'=SUMIFS({RNG("G")},{RNG("F")},"flat*",{RNG("K")},"Y")', 'insul_roof', 1)
L('Insulation', 'Spray foam — vaulted ceilings', 'along the slope', f'=SUMIFS({RNG("G")},{RNG("F")},"vaulted",{RNG("K")},"Y")', 'insul_spray', 1)
L('Drywall', 'Drywall hang + finish — walls', 'finished rooms', f'={FIN("E")}', 'drywall', 1)
L('Drywall', 'Drywall hang + finish — ceilings', 'flat + vaulted', f'={FIN("G")}', 'drywall', 1)
L('Drywall', 'Fire-rated wall, garage side', 'Framing sheet: fire-rated walls', f'={FIRE}', 'firewall', 1)
L('Flooring', 'LVP', 'Rooms: floor finish = LVP', f'=SUMIFS({RNG("D")},{RNG("C")},"LVP")', 'floor_lvp', 1)
L('Flooring', 'Tile floors', 'Rooms: floor finish = tile', f'=SUMIFS({RNG("D")},{RNG("C")},"tile")', 'floor_tile', 1)
L('Flooring', 'Carpet / turf', 'Rooms: floor finish = carpet / turf', f'=SUMIFS({RNG("D")},{RNG("C")},"carpet / turf")', 'turf', 1)
L('Trim & Carpentry', 'Baseboard', 'finished rooms', f'={FIN("J")}', 'base_replace', 1)
L('Trim & Carpentry', 'Door + window casing', '~17 lf per door side, ~16 lf per window',
  f'=(({OC("Interior door", "New")}+{OC("Doorway", "Add door")})*2+{OC("Doorway", "Cased")}*2+{OC("Closet opening", "Bifold")})*17+COUNTIFS(Openings!$B${o0}:$B${oN},"Window")*16', 'casing_replace', 1)
L('Trim & Carpentry', 'Closet shelving + rods', 'closets in the plan', f'=COUNTIFS(Openings!$C${o0}:$C${oN},"*closet*")', 'closet_fit', 1)
if any('pantry' in n.lower() for n in room_names):
    L('Trim & Carpentry', 'Pantry shelving', '', 1, 'pantry_shelf', 1)
L('Paint', 'Paint walls', 'finished rooms', f'={FIN("E")}', 'paint_walls', 1)
L('Paint', 'Paint ceilings', '', f'={FIN("G")}', 'paint_ceiling', 1)
L('Paint', 'Paint trim + doors', 'base + casing', '=D{base}+D{casing}', 'paint_trim', 1)
L('Cabinets & Counters', 'Base cabinets', 'Cabinets sheet', f'={KLF("Base cabinet")}', 'cab_base_only', 1)
L('Cabinets & Counters', 'Upper cabinets', 'Cabinets sheet', f'={KLF("Upper cabinet")}', 'cab_upper', 1)
L('Cabinets & Counters', 'Tall / pantry cabinets', 'Cabinets sheet', f'={KLF("Tall cabinet")}', 'cab_tall', 1)
L('Cabinets & Counters', 'Kitchen island', 'Cabinets sheet', f'={KLF("Island")}', 'cab_island', 1)
L('Cabinets & Counters', 'Quartz countertops', 'base runs + island', f'={QUARTZ}', 'counter', 1)
L('Cabinets & Counters', 'Backsplash', 'kitchen base runs, 18" high', f'=SUMIFS({KR("G")},{KR("H")},"Base cabinet",{KR("A")},"Kitchen*")*1.5', 'backsplash', 1)
for v in van:
    key = 'vanity_72' if v[2] >= 60 else 'vanity_36' if v[2] >= 30 else 'vanity_24'
    L('Cabinets & Counters', f'Vanity {v[2]}" — {v[0]}', v[7], 1, key, 1)
if kitchens:
    L('Cabinets & Counters', 'Appliance package', 'range, hood/micro, fridge, DW', 1, 'appliances', 1)
n_kit = len(kitchens)
n_wash = max(1 if laundry else 0, sum(1 for r in rows if 'washer' in r[1].lower()))
L('Plumbing', 'Rough + set, per fixture', 'toilets, tubs, showers, lavs, kitchen sink + DW, washer',
  f'={KCNT("Toilet")}+{KCNT("Tub")}+{KCNT("Shower")}+{lavs}+{2 * n_kit}+{n_wash}', 'plumb_fixture', 1)
L('Plumbing', 'Toilets', '', f'={KCNT("Toilet")}', 'toilet', 1)
L('Plumbing', 'Tubs', '', f'={KCNT("Tub")}', 'tub', 1)
L('Plumbing', 'Tub / shower valves + trim', '', f'={KCNT("Tub")}+{KCNT("Shower")}', 'shower_valve', 1)
L('Plumbing', 'Lav faucets', 'a 60"+ vanity has two', lavs, 'lav_faucet', 1)
if kitchens:
    L('Plumbing', 'Kitchen sink + faucet + disposal', '', n_kit, 'kitchen_sink', 1)
L('Plumbing', 'Wall tile — tub + shower surrounds', 'tubs 6 ft up, showers 8 ft up', round(tile_wall), 'tile_wall', 1)
if n_show:
    L('Plumbing', 'Shower pans + curbs', '', n_show, 'shower_pan', 1)
    L('Plumbing', 'Shower glass', '', n_show, 'shower_glass', 1)
L('Plumbing', 'Bath accessories + mirrors', 'per bath', len(baths), 'bath_access', 1)
L('Electrical', 'Electrical — finished area', 'wiring, devices, circuits', f'={fin_sf}', 'elec_new', 1)
L('Electrical', 'Light fixtures', '1 per finished room + 4 outside', f'=COUNTIFS({RNG("K")},"Y")+4', 'light_fixture', 1)
if kitchens:
    L('Electrical', 'Recessed lights — kitchen', '1 per ~40 sf', f'=ROUND(SUMIFS({RNG("D")},Rooms!$A${r0}:$A${rN},"Kitchen*")/40,0)', 'recessed', 1)
L('Electrical', 'Smoke / CO detectors', 'each bedroom + hall + living', len(bedrooms) + 2, 'smoke_co', 1)
L('Electrical', 'Bath / laundry exhaust fans', '', len(baths) + (1 if laundry else 0), 'exhaust_fan', 1)
L('Electrical', 'Panel / service upgrade to 200 A', 'verify the existing panel', 1, 'panel_upgrade', 1 if added_sf else 0)
if any(s.get('stretched') for s in t['slabs']):
    L('Electrical', 'Garage electrical (enlarged garage)', 'lights, outlets, door openers', 1, 'electrical_allow', 1)
if added_sf:
    L('HVAC', 'Extend ducts / new zone — added area', 'additions', added_sf, 'hvac_extend', 1)
L('HVAC', 'Service existing system', '', 1, 'hvac_allow', 1)
for e in t['extras']:          # design scope items the model has no geometry for (meter / AC moves, chimney, closed openings...)
    L('Design changes', e['label'], e['note'], e['qty'], e['key'], 1)

b0 = 5
idx = {}
for i, (trade, item, notes, q, key, inc) in enumerate(lines):
    r = b0 + i
    d, u, p, src = rate(key)
    if item == 'Baseboard': idx['base'] = r
    if item == 'Door + window casing': idx['casing'] = r
    put(B, r, 1, trade); put(B, r, 2, item); put(B, r, 3, notes)
    isf = isinstance(q, str) and q.startswith('=')
    put(B, r, 4, q, NUM1 if u in ('sf', 'lf') else NUM, None if isf else BLUE)
    put(B, r, 5, u); put(B, r, 6, p, MONEY2, BLUE); put(B, r, 7, inc, '0', BLUE)
    put(B, r, 8, f'=D{r}*F{r}*G{r}', MONEY); put(B, r, 9, src, color=RED if src.startswith('added') else None)
bN = b0 + len(lines) - 1
for r in range(b0, bN + 1):
    v = B.cell(r, 4).value
    if isinstance(v, str) and '{base}' in v:
        B.cell(r, 4).value = v.replace('{base}', str(idx['base'])).replace('{casing}', str(idx['casing']))
put(B, bN + 2, 7, 'Subtotal', bold=True); put(B, bN + 2, 8, f'=SUM(H{b0}:H{bN})', MONEY, bold=True)
put(B, bN + 4, 1, 'How to use:', bold=True)
put(B, bN + 5, 1, 'Edit blue cells only: Qty overrides, Unit cost, Include (1/0). Black cells are formulas. Red "added — verify" rates are not in the pricebook yet — get a local price.')
TRADES = list(dict.fromkeys(l[0] for l in lines))

# ======================================================================= Summary
U = wb.active; U.title = 'Summary'; U.sheet_view.showGridLines = False
for i, w in enumerate([30, 16, 12, 12, 60]):
    U.column_dimensions[get_column_letter(i + 1)].width = w
U['A1'] = f"{t['address']} — {t['name']} rehab budget"; U['A1'].font = f(bold=True, size=15)
U['A2'] = 'Takeoff from the design model. Installed costs (labor + material).'; U['A2'].font = f(italic=True, color='555555')
header(U, 4, ['Trade', 'Cost', '% of total', '', 'Notes'], freeze=False)
u0 = 5
for i, tr in enumerate(TRADES):
    r = u0 + i
    put(U, r, 1, tr); put(U, r, 2, f'=SUMIFS(Budget!$H${b0}:$H${bN},Budget!$A${b0}:$A${bN},A{r})', MONEY)
    put(U, r, 3, f'=IFERROR(B{r}/$B${u0 + len(TRADES) + 2},0)', '0.0%')
uN = u0 + len(TRADES) - 1
sub = uN + 1
put(U, sub, 1, 'Subtotal', bold=True); put(U, sub, 2, f'=SUM(B{u0}:B{uN})', MONEY, bold=True)
put(U, sub + 1, 1, 'Contingency'); put(U, sub + 1, 4, 0.10, '0.0%', BLUE, fill=YEL); put(U, sub + 1, 2, f'=B{sub}*D{sub+1}', MONEY)
U.cell(sub + 1, 4).comment = Comment('Contingency rate — edit. 10% default.', 'Rebuilt')
tot = sub + 2
put(U, tot, 1, 'Total budget', bold=True, fill=SUB); put(U, tot, 2, f'=B{sub}+B{sub+1}', MONEY, bold=True, fill=SUB)
put(U, tot + 2, 1, 'Finished floor area (sf)'); put(U, tot + 2, 2, f'=Rooms!D{RT}', NUM, GREEN)
put(U, tot + 3, 1, 'Cost per finished sf'); put(U, tot + 3, 2, f'=IFERROR(B{tot}/B{tot+2},0)', MONEY)
placeholder = sum(1 for _, _, _, _, k, _ in lines if k in PB and PB[k].get('source') == 'placeholder')
notes = ['Every quantity comes from the 3D model of the design (field-verify before ordering).',
         f'Unit costs: {placeholder} lines use Rebuilt pricebook rates still marked placeholder; red "added — verify" rates are new. Replace both with GC / sub numbers.' if placeholder else 'Unit costs from the Rebuilt pricebook; red "added — verify" rates are new.',
         'Allowances (red on the Cabinets tab) cover things the design needs but has not drawn yet.',
         'Structure assumes rafters unless the roof is set to trusses in the editor. Beam, header, rafter and footing sizes are a first pass (20 psf snow, 1500 psf soil) — an engineer confirms them.',
         'Material order counts are on the Material List tab — add supplier prices there for a material-only number.',
         'Edit blue cells. Black cells are formulas. Green cells link to another tab.']
put(U, tot + 6, 1, 'Read me', bold=True)
for i, s in enumerate(notes):
    put(U, tot + 7 + i, 1, f'• {s}')

# ======================================================================= Material list
Mt = sheet('Material List', 'Material order list', 'Order quantities with waste. Add supplier prices in the yellow column for a material-only total. Row 5 has an example price — replace it.',
           [18, 48, 12, 7, 12, 8, 11, 12, 13, 13])
header(Mt, 4, ['Group', 'Item', 'Takeoff qty', 'Unit', 'Per pack / coverage', 'Waste', 'Order qty', 'Order unit', 'Supplier price', 'Extended'])
mats = [('Framing', '2x4 studs (partitions)', f"={FR('I', '2x4 partition')}", 'ea', 1, 0.10, 'ea'),
        ('Framing', '2x6 studs (outside walls)', f"={FR('I', '2x6 exterior')}", 'ea', 1, 0.10, 'ea'),
        ('Framing', '2x4 plates', f"={FR('J', '2x4 partition')}", 'lf', 16, 0.10, '16 ft'),
        ('Framing', '2x6 plates', f"={FR('J', '2x6 exterior')}", 'lf', 16, 0.10, '16 ft'),
        ('Framing', '7/16" OSB wall sheathing', f"={FR('K', '2x6 exterior')}", 'sheet', 1, 0.0, 'sheet'),
        ('Framing', 'House wrap (9 x 150 ft roll)', f"={FR('F', '2x6 exterior')}", 'sf', 1350, 0.10, 'roll')]
for i, (nm, mem, n, Lm) in enumerate(rfr):
    r = ra + i
    if 'truss' in mem.lower():
        mats.append(('Framing', f'Trusses — {nm[:40]} ({Lm} ft span)', f"='Framing & Structure'!E{r}", 'ea', 1, 0.0, 'ea'))
    else:
        mats.append(('Framing', f'{mem.split()[0]} rafters — {nm[:40]}', f"='Framing & Structure'!E{r}*'Framing & Structure'!F{r}", 'lf', 16, 0.05, '16 ft'))
mats += [('Framing', '7/16" OSB roof sheathing — new roofs', f'={ROOF_NEW}', 'sf', 32, 0.10, 'sheet')]
for dep in sorted({b['depth_in'] for b in beams}):
    rr = [ba + i for i, b in enumerate(beams) if b['depth_in'] == dep]
    mats.append(('Structure', f'LVL 1-3/4" x {dep:g}" (plies x length)', '=' + '+'.join(f"'Framing & Structure'!E{r}*'Framing & Structure'!H{r}" for r in rr), 'lf', 1, 0.05, 'lf'))
npost = sum(2 * b['count'] if not b.get('ridge') else b['count'] + 1 for b in beams)
if beams:
    mats.append(('Structure', 'Post footings (concrete pads under beam posts, see Framing tab)', f'={npost}', 'ea', 1, 0.0, 'ea'))
mats += [('Roofing', 'Architectural shingles (3 bundles / square)', f"='Roof & Exterior'!D{xN+1}", 'sq', 1 / 3, 0.10, 'bundle'),
         ('Roofing', 'Ice & water shield (2 sq / roll)', f"={EDGE['iw1']}+{EDGE['iw2']}", 'sf', 200, 0.10, 'roll'),
         ('Roofing', 'Synthetic underlayment (10 sq / roll)', f"={EDGE['syn']}", 'sf', 1000, 0.10, 'roll'),
         ('Roofing', 'Drip edge (10 ft)', f"={EDGE['drip']}", 'lf', 10, 0.05, 'piece'),
         ('Roofing', 'Ridge cap (~25 lf / bundle)', f"={EDGE['ridge']}", 'lf', 25, 0.05, 'bundle'),
         ('Roofing', 'Starter strip (~100 lf / bundle)', f"='Roof & Exterior'!E{xN+1}", 'lf', 100, 0.05, 'bundle'),
         ('Exterior', 'Vinyl siding (2 sq / box)', f"={EDGE['sid']}", 'sf', 200, 0.10, 'box'),
         ('Exterior', 'Gutter (seamless, by the foot)', f"={EDGE['gut']}", 'lf', 1, 0.05, 'lf'),
         ('Exterior', 'Downspouts (10 ft)', f"={EDGE['ds']}", 'lf', 10, 0.0, 'piece'),
         ('Insulation', 'R-20 wall batts (~40 sf / bag)', f'={INS_EXT}', 'sf', 40, 0.05, 'bag'),
         ('Insulation', 'Blown R-49 (~40 sf / bag)', f'=SUMIFS({RNG("G")},{RNG("F")},"flat*",{RNG("K")},"Y")', 'sf', 40, 0.05, 'bag'),
         ('Drywall', '1/2" drywall 4x12 — walls', f'={FIN("E")}', 'sf', 48, 0.12, 'sheet'),
         ('Drywall', '1/2" ceiling board 4x12 — ceilings', f'={FIN("G")}', 'sf', 48, 0.12, 'sheet'),
         ('Drywall', '5/8" Type X 4x12 — garage side', f'={FIRE}', 'sf', 48, 0.12, 'sheet'),
         ('Drywall', 'Joint compound (4.5 gal box, ~450 sf)', f'={FIN("E")}+{FIN("G")}', 'sf', 450, 0.0, 'box'),
         ('Flooring', 'LVP (~20 sf / box)', f'=SUMIFS({RNG("D")},{RNG("C")},"LVP")', 'sf', 20, 0.10, 'box'),
         ('Flooring', 'Floor tile', f'=SUMIFS({RNG("D")},{RNG("C")},"tile")', 'sf', 1, 0.12, 'sf'),
         ('Flooring', 'Wall tile (surrounds + showers)', round(tile_wall), 'sf', 1, 0.12, 'sf'),
         ('Flooring', 'Cement backer board 3x5', f'=SUMIFS({RNG("D")},{RNG("C")},"tile")+{round(tile_wall)}', 'sf', 15, 0.10, 'sheet'),
         ('Flooring', 'Thinset (50 lb, ~50 sf)', f'=SUMIFS({RNG("D")},{RNG("C")},"tile")+{round(tile_wall)}', 'sf', 50, 0.10, 'bag'),
         ('Trim', 'Baseboard (16 ft)', f'={FIN("J")}', 'lf', 16, 0.10, 'piece'),
         ('Trim', 'Casing (7 ft)', f'=Budget!D{idx["casing"]}', 'lf', 7, 0.10, 'piece'),
         ('Paint', 'Wall + ceiling paint, 2 coats (~350 sf / gal)', f'=({FIN("E")}+{FIN("G")})*2', 'sf', 350, 0.05, 'gal'),
         ('Paint', 'Primer, 1 coat (~300 sf / gal)', f'={FIN("E")}+{FIN("G")}', 'sf', 300, 0.05, 'gal')]
m0 = 5
for i, (grp, item, q, u, per, waste, ou) in enumerate(mats):
    r = m0 + i
    isf = isinstance(q, str) and q.startswith('=')
    put(Mt, r, 1, grp); put(Mt, r, 2, item); put(Mt, r, 3, q, NUM1, None if isf else BLUE); put(Mt, r, 4, u)
    put(Mt, r, 5, per, '0.###', BLUE); put(Mt, r, 6, waste, '0%', BLUE)
    put(Mt, r, 7, f'=ROUNDUP(C{r}*(1+F{r})/E{r},0)', NUM); put(Mt, r, 8, ou)
    put(Mt, r, 9, None, MONEY2, BLUE, fill=YEL); put(Mt, r, 10, f'=IF(I{r}="","",G{r}*I{r})', MONEY)
Mt.cell(m0, 9).value = 4.25
Mt.cell(m0, 9).comment = Comment('Example only — replace with your supplier quote (per stud).', 'Rebuilt')
mN = m0 + len(mats) - 1
put(Mt, mN + 2, 1, 'Windows, doors, cabinets + fixtures: see the Openings and Cabinets & Fixtures tabs for the full schedule with sizes.').font = f(italic=True, color='555555')
put(Mt, mN + 3, 9, 'Material total', bold=True); put(Mt, mN + 3, 10, f'=SUM(J{m0}:J{mN})', MONEY, bold=True)

wb.move_sheet('Summary', -10)
wb.save(A.out)
print('wrote', A.out, '|', len(lines), 'budget lines')
